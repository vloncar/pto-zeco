#!/usr/bin/env python3
"""Split pto_einsum_slide.pptx into two slides and add the GDN result.

    slide 1  compiler and layout      PTO-EINSUM / PTO-FUSER, grown to fill the
                                      space the moved half leaves behind
    slide 2  pypto                    ZECO / ALLSCAN on top, GATED DELTANET below

Always starts from pto_einsum_slide.1slide.pptx, the single-slide deck as it
stood before this script, so re-running is idempotent and never stacks edits.
Copy follows the deck's own pattern: a bold lead claim in 1D1D1A followed by a
lighter 3F3F3D elaboration, and the deck's vocabulary ("generated code" /
"hand-written code") rather than internal tool names.
"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches, Pt

ROOT = Path("/root/workspace/allscan")
DECK = ROOT / "pto_einsum_slide.pptx"
BASE = ROOT / "pto_einsum_slide.1slide.pptx"
GDN_CHART = ROOT / "devtools" / "slide_plot_gdn.png"
GDN_NOTES = ROOT / "devtools" / "slide_notes_gdn.txt"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# Shape ids in the base deck.
EINSUM_HEADER = (3, 4, 5)            # accent square, section label, title + subtitle
EINSUM_BULLETS = 6
EINSUM_CHART = (9, 10, 11, 12, 13, 14) + tuple(range(15, 49)) + (49,)
ZECO = (51, 52, 53, 54, 55, 56)      # accent, label, title, left col, right col, caption
PANEL, DIVIDER = 2, 50

# --- slide 2 geometry (inches) ----------------------------------------------
Z_ACCENT_T, Z_LABEL_T, Z_TITLE_T, Z_BULLETS_T = 0.52, 0.44, 0.66, 1.50
Z_TITLE_H, Z_BULLETS_H, Z_BULLET_FONT = 0.80, 2.30, 10.5
COL_L, COL_R, COL_HALF = 0.80, 6.94, 5.59

G_ACCENT_T, G_LABEL_T, G_TITLE_T, G_BULLETS_T = 4.17, 4.10, 4.32, 4.95
G_TITLE_H, G_BULLETS_H, G_TEXT_W = 0.50, 1.85, 6.00
G_BULLET_FONT = 9.5
G_CHART_T, G_CHART_H = 4.36, 2.24
G_CAPTION_T, G_CAPTION_H = 6.62, 0.30

GDN_LABEL = "GATED DELTANET"
GDN_TITLE = "Hand-tuned chip kernels, rebuilt by the compiler"
GDN_SUBTITLE = "matched stage for stage, and improved on in two places"
GDN_BULLETS = [
    ("The attention layer inside Qwen3.5 and 3.6.",
     " Gated DeltaNet stands in for most of the attention in those models, so its speed "
     "sets theirs."),
    ("Six hand-tuned kernels, rewritten in the compiler's own language.",
     " Same algorithm, same six stages, same accuracy target as the shipping hand-written "
     "code — which is what makes the comparison fair."),
    ("Feature parity, and two improvements to hand back.",
     " Every stage agrees with the reference to the precision the hardware allows, and end "
     "to end the generated code is slightly quicker. The rewrite also turned up two "
     "shortcuts the hand-written version had not taken."),
]
GDN_CAPTION = ("The hand-written version's own benchmark, run unchanged: 8,192 tokens, 16 heads, "
               "width 128. Median of 50 runs on one Ascend 910B2.")

# --- slide 1 geometry (inches), grown into the freed half --------------------
E_LABEL_T, E_ACCENT_T = 0.52, 0.60
E_TITLE_T, E_TITLE_H = 0.78, 1.55
E_BULLETS_T, E_BULLETS_H, E_BULLETS_W = 2.50, 4.10, 6.45
E_CHART_ANCHOR, E_CHART_TOP, E_CHART_SCALE = 0.46, 1.00, 1.686
E_BAR_SCALE = 2.0                     # bar thickness grows faster than the spacing
E_FONTS = {4: 11.0, 9: 11.0, 11: 8.5, 13: 8.5, 49: 9.0}
E_HEADER_L, E_HEADER_W = 7.25, 3.98    # one line at the grown font size
E_TITLE_FONTS = (21.0, 16.0)          # title, subtitle
E_BULLET_FONT = 12.0
E_ROW_FONT, E_AXIS_FONT = 9.0, 8.0
#: The chart labels grew, so their boxes have to. Row labels are right-aligned and
#: axis labels centred, so both widen away from the bars rather than into them.
E_ROW_LABEL_L, E_ROW_LABEL_W = 7.25, 1.58
E_AXIS_LABEL_DX, E_AXIS_LABEL_W = -0.13, 1.10
E_RATIO_W = 0.55


def _shapes(slide):
    return {sh.shape_id: sh for sh in slide.shapes}


def _drop(shape):
    shape._element.getparent().remove(shape._element)


def _place(shape, left=None, top=None, width=None, height=None):
    if left is not None:
        shape.left = Inches(left)
    if top is not None:
        shape.top = Inches(top)
    if width is not None:
        shape.width = Inches(width)
    if height is not None:
        shape.height = Inches(height)


def _set_font(shape, size, para_index=None):
    for i, para in enumerate(shape.text_frame.paragraphs):
        if para_index is not None and i != para_index:
            continue
        for run in para.runs:
            run.font.size = Pt(size)


def _restyle(run, bold, colour, text):
    rPr = run.find(f"{A}rPr")
    rPr.set("b", "1" if bold else "0")
    rPr.find(f"{A}solidFill").find(f"{A}srgbClr").set("val", colour)
    run.find(f"{A}t").text = text


def _set_bullets(shape, bullets):
    """Rewrite a bulleted box as bold lead plus elaboration, cloning its formatting."""
    body = shape.text_frame._txBody
    paras = body.findall(f"{A}p")
    first, rest_tpl = paras[0], (paras[1] if len(paras) > 1 else paras[0])
    for p in paras:
        body.remove(p)
    for i, (lead, tail) in enumerate(bullets):
        p = copy.deepcopy(first if i == 0 else rest_tpl)
        runs = p.findall(f"{A}r")
        for extra in runs[1:]:
            p.remove(extra)
        _restyle(runs[0], True, "1D1D1A", lead)
        tail_run = copy.deepcopy(runs[0])
        _restyle(tail_run, False, "3F3F3D", tail)
        runs[0].addnext(tail_run)
        body.append(p)


def _set_run(shape, text, para_index=0):
    p = shape.text_frame._txBody.findall(f"{A}p")[para_index]
    p.find(f"{A}r").find(f"{A}t").text = text


def _clone_into(slide, shape):
    """Deep-copy a shape's XML into *slide* and return the new shape object."""
    element = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(element)
    return slide.shapes[-1]


def build_slide_two(prs, source):
    """New slide: ZECO / ALLSCAN moved to the top half, GATED DELTANET below."""
    slide = prs.slides.add_slide(source.slide_layout)
    for ph in list(slide.shapes):                 # the layout seeds empty placeholders
        _drop(ph)

    src = _shapes(source)
    panel = _clone_into(slide, src[PANEL])        # grey ground for the bottom half
    _clone_into(slide, src[DIVIDER])
    moved = {sid: _clone_into(slide, src[sid]) for sid in ZECO}

    _place(moved[51], COL_L, Z_ACCENT_T)
    _place(moved[52], 1.04, Z_LABEL_T)
    _place(moved[53], COL_L, Z_TITLE_T, 11.73, Z_TITLE_H)
    _place(moved[54], COL_L, Z_BULLETS_T, COL_HALF, Z_BULLETS_H)
    _place(moved[55], COL_R, Z_BULLETS_T, COL_HALF, Z_BULLETS_H)
    _set_font(moved[54], Z_BULLET_FONT)
    _set_font(moved[55], Z_BULLET_FONT)
    _place(moved[56], 10.81, 0.40)

    accent = _clone_into(slide, moved[51])
    _place(accent, COL_L, G_ACCENT_T)
    label = _clone_into(slide, moved[52])
    _place(label, 1.04, G_LABEL_T, 5.50)
    _set_run(label, GDN_LABEL)
    title = _clone_into(slide, moved[53])
    _place(title, COL_L, G_TITLE_T, G_TEXT_W, G_TITLE_H)
    _set_run(title, GDN_TITLE, 0)
    _set_run(title, GDN_SUBTITLE, 1)
    _set_font(title, 13.5, para_index=0)
    _set_font(title, 11.0, para_index=1)
    bullets = _clone_into(slide, moved[54])
    _place(bullets, COL_L, G_BULLETS_T, G_TEXT_W, G_BULLETS_H)
    _set_bullets(bullets, GDN_BULLETS)
    _set_font(bullets, G_BULLET_FONT)

    slide.shapes.add_picture(str(GDN_CHART), Inches(COL_R), Inches(G_CHART_T),
                             Inches(COL_HALF), Inches(G_CHART_H))
    caption = _clone_into(slide, src[49])         # the einsum footnote's exact styling
    _place(caption, COL_R, G_CAPTION_T, COL_HALF, G_CAPTION_H)
    _set_run(caption, GDN_CAPTION)

    for sid in ZECO:
        _drop(src[sid])
    for sid in (PANEL, DIVIDER):
        _drop(src[sid])
    del panel
    return slide


def grow_slide_one(slide):
    """With the bottom half gone, let PTO-EINSUM / PTO-FUSER use the whole page."""
    by = _shapes(slide)

    _place(by[3], top=E_ACCENT_T)
    _place(by[4], top=E_LABEL_T)
    _place(by[5], top=E_TITLE_T, height=E_TITLE_H)
    _set_font(by[5], E_TITLE_FONTS[0], para_index=0)
    _set_font(by[5], E_TITLE_FONTS[1], para_index=1)

    _place(by[6], top=E_BULLETS_T, width=E_BULLETS_W, height=E_BULLETS_H)
    _set_font(by[6], E_BULLET_FONT)

    for sid in EINSUM_CHART:
        sh = by[sid]
        old_top = Emu(sh.top).inches
        old_h = Emu(sh.height).inches
        grow = E_BAR_SCALE if sh.name.startswith("Rectangle") and old_h < 0.2 else E_CHART_SCALE
        _place(sh,
               top=E_CHART_TOP + (old_top - E_CHART_ANCHOR) * E_CHART_SCALE,
               height=old_h * grow)
        if sid in E_FONTS:
            _set_font(sh, E_FONTS[sid])
            if sid == 9:
                _place(sh, left=E_HEADER_L, width=E_HEADER_W)
        elif 15 <= sid <= 44 and sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            # the bars in this range are auto shapes; their widths carry the values
            _set_font(sh, E_ROW_FONT)
            if sh.text_frame.paragraphs[0].alignment is not None:      # right-aligned row label
                _place(sh, left=E_ROW_LABEL_L, width=E_ROW_LABEL_W)
            else:                                                      # ratio label
                _place(sh, width=E_RATIO_W)
        elif 45 <= sid <= 48:
            _set_font(sh, E_AXIS_FONT)
            _place(sh, left=Emu(sh.left).inches + E_AXIS_LABEL_DX, width=E_AXIS_LABEL_W)

    # The axis rule spans the rows, so it takes the row scale, not the bar scale.
    _place(by[14], height=Emu(by[14].height).inches / E_BAR_SCALE * E_CHART_SCALE)


def main() -> int:
    if not BASE.exists():
        raise SystemExit(f"missing {BASE} — cannot guarantee a clean base")
    if not GDN_CHART.exists():
        raise SystemExit(f"missing {GDN_CHART} — run make_gdn_slide_plot.py first")
    shutil.copy2(BASE, DECK)

    prs = Presentation(DECK)
    first = prs.slides[0]
    second = build_slide_two(prs, first)
    grow_slide_one(first)

    # The speaker notes belong to the halves, and one half has moved.
    zeco_notes = ""
    if first.has_notes_slide:
        zeco_notes = first.notes_slide.notes_text_frame.text.replace(
            "ZECO / ALLSCAN — bottom half.", "ZECO / ALLSCAN — top half.")
        first.notes_slide.notes_text_frame.text = ""
    gdn_notes = GDN_NOTES.read_text() if GDN_NOTES.exists() else ""
    second.notes_slide.notes_text_frame.text = (
        zeco_notes.rstrip() + "\n\n" + "=" * 78 + "\n\n" + gdn_notes).strip()

    prs.save(DECK)
    print(f"saved {DECK}: {len(prs.slides)} slides")
    for i, s in enumerate(prs.slides):
        print(f"  slide {i + 1}: {len(s.shapes)} shapes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
